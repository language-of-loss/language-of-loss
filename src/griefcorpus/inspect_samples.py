import collections
import io
import json
import re
import warnings
import zipfile

from .http import PoliteClient, ROOT, now_iso

warnings.filterwarnings("ignore")
SAMPLES = ROOT / "data" / "discovery" / "samples"
OUT = ROOT / "data" / "discovery" / "evidence" / "file_inspection.json"
WORD = re.compile(r"[^\W\d_]+")
DEV = re.compile(r"[ऀ-ॿ]")

TARGETS = [
    dict(key="dataverse::doi:10.7910/DVN/QDENSA", kind="docx", purge_after=True, files={
        "11033215.docx": "https://dataverse.harvard.edu/api/access/datafile/11033215",
        "11033216.docx": "https://dataverse.harvard.edu/api/access/datafile/11033216",
        "11033217.docx": "https://dataverse.harvard.edu/api/access/datafile/11033217"}),
    dict(key="zenodo::15591374", kind="docx", files={
        "widows.docx": "https://zenodo.org/api/records/15591374/files/Widows%20interview%20as%20a%20narrative%20for%20analysis.docx/content"}),
    dict(key="zenodo::22041925", kind="hinemo_zip", files={
        "hinemo9.zip": "https://zenodo.org/api/records/22041925/files/im-rakeshsahani/hinemo9-v1.0.1.zip/content"}),
    dict(key="url::https://reshare.ukdataservice.ac.uk/856067/", kind="york_xlsx", files={
        "survey.xlsx": "https://reshare.ukdataservice.ac.uk/856067/2/ExperiencesOfGrief_APhenomenologicalSurvey.xlsx"}),
    dict(key="url::https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0240270", kind="pdf", files={
        "s002.pdf": "https://journals.plos.org/plosone/article/file?type=supplementary&id=10.1371/journal.pone.0240270.s002"}),
]

KIN_EN = r"(my|our) (late )?(mother|mom|mum|father|dad|papa|brother|sister|husband|wife|son|daughter|grand(ma|pa|mother|father)|nani|dadi|nana|dada|friend|uncle|aunt)"
PAT = {
    "en": re.compile(KIN_EN + r"[^.]{0,60}(passed away|died|is no more|expired|left us|death)|lost my (mother|mom|father|dad|papa|brother|sister|husband|wife|son|daughter|grand\w+)|(passed away|died)[^.]{0,40}" + KIN_EN, re.I),
    "hi": re.compile(r"(मेरे|मेरी|मेरा|हमारे|हमारी)[^।]{0,40}(पापा|पिता|पिताजी|माँ|मां|मम्मी|भाई|बहन|पति|पत्नी|बेटा|बेटी|दादी|दादा|नानी|नाना|दोस्त)[^।]{0,60}(नहीं रहे|नहीं रहीं|गुजर|गुज़र|देहांत|निधन|चल बसे|चले गए|छोड़कर चले|मृत्यु|स्वर्गवास|खो दिया)"),
    "hi-Latn": re.compile(r"(mere|meri|mera|hamare|hamari)[^.]{0,40}(papa|pita|maa|mummy|mom|bhai|behen|behan|pati|patni|beta|beti|dadi|dada|nani|nana|dost)[^.]{0,60}(nahi rahe|nahin rahe|guzar|gujar|dehant|nidhan|chal base|chale gaye|chhod kar|death ho|expire|khatam ho gaye|is duniya me nahi)", re.I),
}


def text_stats(txt):
    letters = sum(ch.isalpha() for ch in txt)
    return {"word_tokens": len(WORD.findall(txt)), "devanagari_share_of_letters": round(len(DEV.findall(txt)) / max(letters, 1), 3)}


def profile(kind, path):
    if kind == "docx":
        import docx
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        return {"paragraphs": len(paras), **text_stats("\n".join(paras))}
    if kind == "pdf":
        from pypdf import PdfReader
        r = PdfReader(str(path))
        head = "".join((p.extract_text() or "") for p in r.pages[:40])
        return {"pages": len(r.pages), "first_40_pages": text_stats(head)}
    if kind == "york_xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(str(path), read_only=True)["Sheet 1 - Bereavement"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        nat = hdr.index("Respondent's Nationality")
        qcols = [i for i, h in enumerate(hdr) if re.match(r"[AB]\d\d\.", h)]
        from .textutil import word_count
        per = [sum(word_count(r[i]) for i in qcols if isinstance(r[i], str)) for r in rows[1:] if r[0] is not None]
        nats = collections.Counter(str(r[nat]).strip().lower() for r in rows[1:] if r[0] is not None)
        return {"respondents": len(per), "word_tokens_in_answer_columns": sum(per),
                "median_word_tokens_per_respondent": sorted(per)[len(per) // 2],
                "nationality_top": nats.most_common(5), "nationality_containing_india": sum(v for k, v in nats.items() if "india" in k)}
    if kind == "hinemo_zip":
        z = zipfile.ZipFile(str(path))
        name = next(n for n in z.namelist() if n.endswith("data/full.jsonl"))
        recs = [json.loads(l) for l in io.TextIOWrapper(z.open(name), encoding="utf-8")]
        hits = collections.Counter()
        for r in recs:
            if any(p.search(r["text"]) for p in PAT.values()):
                hits[r["language_label"]] += 1
        return {"records": len(recs), "by_variety": collections.Counter(r["language_label"] for r in recs),
                "by_label_source": collections.Counter(r["label_source"] for r in recs),
                "fields": sorted(recs[0]), "narrow_first_person_bereavement_hits": hits,
                "hit_rate": round(sum(hits.values()) / len(recs), 5)}


def run():
    client = PoliteClient(min_delay_s=1.5, timeout_s=120)
    report = {"inspected_at": now_iso(), "sources": {}}
    for t in TARGETS:
        folder = SAMPLES / re.sub(r"[^A-Za-z0-9]+", "_", t["key"])[:60]
        folder.mkdir(parents=True, exist_ok=True)
        entry = {"files": {}}
        for fname, url in t["files"].items():
            path = folder / fname
            if not path.exists():
                resp, body = client.get(url)
                if resp is None or resp.status_code != 200:
                    entry["files"][fname] = {"http_status": getattr(resp, "status_code", None), "downloaded": False}
                    continue
                path.write_bytes(body)
            entry["files"][fname] = {"bytes": path.stat().st_size, "profile": profile(t["kind"], path)}
            if t.get("purge_after"):
                path.unlink()
                entry["files"][fname]["purged_after_profiling"] = True
        report["sources"][t["key"]] = entry
        print(t["key"], json.dumps(entry, ensure_ascii=False, default=dict)[:400], flush=True)
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=1, default=dict), encoding="utf-8")


if __name__ == "__main__":
    run()
