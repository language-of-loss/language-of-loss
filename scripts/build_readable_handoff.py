import csv
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from griefcorpus.textutil import word_count

OUT = ROOT / "release" / "For collaborator - Grief project data v1"
BUILT = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
CUTOFF = "2026-09-19 14:15 UTC"
PART = 32000
ARIAL, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE6F0")
LANG = {"english": "English", "devanagari_hindi": "Hindi (Devanagari)", "romanized_hindi": "Hindi (Romanized)",
        "mixed_hi_en": "Mixed Hindi-English", "uncertain": "Uncertain"}
LANG_ORDER = list(LANG.values())


FLAG_WORDS = {"over_cap": "extra page from the same site or author", "length_outlier_gt5000w": "very long (over 5,000 words)",
              "focus_early": "death mentioned early", "provenance:repost": "repost of an earlier piece", "provenance:unclear": "unclear whether original"}


def plain(v):
    if not isinstance(v, str):
        return v
    v = re.sub(r"boilerplate_lines_removed:(\d+)", r"\1 lines of site text removed from the end", v)
    for k, w in FLAG_WORDS.items():
        v = v.replace(k, w)
    for k, w in LANG.items():
        v = v.replace(k, w)
    v = v.replace("hi-Latn", "Romanized Hindi").replace("hi_en", "Hindi-English")
    return v.replace("regex screen", "keyword screen").replace("`", "").replace("_", " ")


def unmark(line):
    line = re.sub(r"`?r([12])_reviewer_id`?", r"'Reviewer \1: reviewer id'", line)
    line = re.sub(r"`?r([12])_`?", r"'Reviewer \1'", line)
    line = line.replace("`adjudicated_`", "'Agreed after discussion'").replace("`annotation_id`", "the ID in the first column")
    line = re.sub(r"^#+\s*", "", line).replace("**", "").replace("*", "")
    return plain(line)


def jl(name):
    return [json.loads(l) for l in (ROOT / "data" / "clean" / f"{name}.jsonl").open(encoding="utf-8")]


def parts(text, n):
    text = ILLEGAL_CHARACTERS_RE.sub("", text or "")
    chunks = [text[i:i + PART] for i in range(0, len(text), PART)] or [""]
    return chunks + [""] * (n - len(chunks))


def sheet(wb, title, header, rows, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name="Arial", size=10, italic=True)
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill = BOLD, HEAD_FILL
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"
            c.font = ARIAL
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws


def text_sheet(wb, title, lines):
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 130
    for i, line in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.font = BOLD if line.startswith("#") or line.isupper() else ARIAL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def new_wb():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def lang_row(label, recs, key="lang_provisional"):
    c = Counter(LANG.get(r[key], r[key]) for r in recs)
    return [label] + [c.get(l, 0) for l in LANG_ORDER] + [len(recs)]


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "DO NOT OPEN until annotation is finished").mkdir(parents=True)
    blogs, ytcv, hin, red, york = jl("blogs"), jl("ytcv"), jl("hinemo9"), jl("reddit_confessions"), jl("york_survey")
    seed = list(csv.DictReader((ROOT / "docs" / "seed_corpus.csv").open(encoding="utf-8")))
    a_sheet = list(csv.DictReader((ROOT / "docs" / "annotation_sheet.csv").open(encoding="utf-8")))
    a_key = list(csv.DictReader((ROOT / "data" / "annotation" / "annotation_key.csv").open(encoding="utf-8")))
    rec = {r["record_id"]: r for src in (blogs, ytcv, hin, red, york) for r in src}
    el = [r for r in blogs if r["eligibility_status"] == "provisionally_eligible"]
    status = lambda r: "candidate personal account (not human-checked)" if r["eligibility_status"] == "provisionally_eligible" \
        else "kept separately: " + r["eligibility_status"].replace("secondary_", "").replace("_", " ")
    caution = ("Every language, type and status label in this workbook was produced by rules or by machine reading. "
               "No human has checked any row. This is not a validated or representative corpus.")

    wb = new_wb()
    text_sheet(wb, "READ ME", [
        "BLOGS AND YOUTUBE COMMENTS - CLEANED DATA (release v1)", f"Built {BUILT}. Data cutoff {CUTOFF}.", "", caution, "",
        "SHEETS",
        "Overview - counts by source and language.",
        f"Blogs - {len(blogs)} web pages: {len(el)} candidate first-person bereavement pieces + {len(blogs) - len(el)} contrast pages (poems, tributes, advice, templates, translations, obituaries). PRIMARY material. Full texts are in workbook 5.",
        f"YouTube comments - {len(ytcv) + len(hin)} comments that passed a bereavement keyword screen, from two open research datasets (not collected from YouTube directly). SECONDARY material.",
        "Seed 15 examples - the 15 feasibility examples read in full by machine; kept out of the annotation queue.",
        "Column guide - what each column means.", "",
        "IMPORTANT LIMITS",
        "Blogs: Hindi pages are mostly 2007-2017 Blogspot death-anniversary posts and literary memoirs; English pages are mostly 2015-2026 news-site personal essays. Language is tangled up with genre and period.",
        "YouTube: the dataset hides video titles, has no dates and no reply links. Nearly all Hindi comments are in Roman script and come from a few song videos. It cannot show how Hindi and English writers respond to the same video.",
        "About half of the YouTube rows read by machine were real first-person losses; the rest include TV-serial plot talk. Treat 'candidate' as 'worth a human look', nothing more.", "",
        "CREDIT (required when reusing)",
        "YTCommentVerse, Zenodo doi:10.5281/zenodo.15678816, CC BY 4.0. HinEmo-9, Zenodo doi:10.5281/zenodo.22041925, CC BY 4.0.",
        "Blog pages belong to their authors and publishers. Use them for our research only; do not post or forward the texts."])
    over = [lang_row("Blogs - candidate personal accounts", el), lang_row("Blogs - contrast pages", [r for r in blogs if r not in el]),
            lang_row("YouTube - YTCommentVerse", ytcv), lang_row("YouTube - HinEmo-9", hin)]
    sheet(wb, "Overview", ["Source"] + LANG_ORDER + ["Total"], over, [42, 10, 18, 18, 20, 11, 9],
          note=f"Counted from the data sheets on {BUILT} (static values). To re-check, filter the 'language (machine)' column of each sheet. Machine labels, not human judgements.")
    ws = wb["Overview"]
    base = ws.max_row + 2
    extra = [("Blog words - Hindi (Devanagari), candidates", sum(r["n_words"] for r in el if r["lang_provisional"] == "devanagari_hindi")),
             ("Blog words - English, candidates", sum(r["n_words"] for r in el if r["lang_provisional"] == "english")),
             ("Blog candidates over 5,000 words (all Hindi memoirs)", sum("length_outlier_gt5000w" in r["flags"] for r in el)),
             ("Distinct hashed bylines among blog candidates", len({r["author_group"] for r in el if r["author_group"]})),
             ("Blog candidates with no byline on the page", sum(1 for r in el if not r["author_group"])),
             ("YouTube: distinct (hashed) videos", len({r["thread_id"] for r in ytcv})),
             ("YouTube: share of the source dataset scanned", "22.2% (download was cut off by the server; frozen there)"),
             ("Rows checked by a human", 0)]
    for i, (k, v) in enumerate(extra):
        ws.cell(row=base + i, column=1, value=k).font = ARIAL
        ws.cell(row=base + i, column=2, value=v).font = ARIAL
    bh = ["record ID", "web address", "site", "title", "published", "language (machine)", "words", "page type (machine-assigned)", "status",
          "who died (as stated on page)", "stated link to India (English list only)", "author group (scrambled, no names)", "remarks", "searcher's note", "first 300 characters"]
    brow = [[r["record_id"], r["source_url"], r["thread_id"], (r.get("extra") or {}).get("title"), r["published_at"], LANG.get(r["lang_provisional"]),
             r["n_words"], r["message_type_provisional"].replace("_", " "), status(r), (r.get("extra") or {}).get("deceased_relation_as_stated"),
             (r.get("extra") or {}).get("indian_context_evidence"), r["author_group"], plain("; ".join(r["flags"])),
             plain((r.get("extra") or {}).get("searcher_notes")), r["text_clean"][:300]] for r in blogs]
    sheet(wb, "Blogs", bh, brow, [22, 45, 24, 40, 12, 18, 8, 22, 34, 22, 34, 18, 26, 50, 70])
    yh = ["record ID", "dataset", "language (machine)", "words", "comment text (cleaned, names/links masked)", "video (scrambled ID)", "commenter group (scrambled)",
          "video category", "votes", "status"]
    yrow = [[r["record_id"], "YTCommentVerse" if r["record_id"].startswith("ytcv") else "HinEmo-9", LANG.get(r["lang_provisional"]), r["n_words"],
             r["text_clean"], r["thread_id"], r["author_group"], (r.get("extra") or {}).get("video_category"), (r.get("extra") or {}).get("votes"),
             "passed keyword screen (not human-checked)"] for r in ytcv + hin]
    sheet(wb, "YouTube comments", yh, yrow, [40, 16, 18, 8, 110, 34, 20, 14, 8, 34])
    sh = ["seed id", "kind of source", "language (machine-reviewed)", "words", "message type (machine-reviewed)", "suitability (machine-reviewed)", "who died", "source", "source record", "machine note", "excerpt"]
    srow = [[plain(s["seed_id"]), plain(s["platform_genre"]), plain(s["language_machine_reviewed"]), int(s["n_words"]), plain(s["message_type_machine_reviewed"]),
             plain(s["eligibility_machine_reviewed"]), s["deceased_relation_as_stated"], s["source_url"], s["source_record_id"], plain(s["machine_note"]), s["text_excerpt"]] for s in seed]
    sheet(wb, "Seed 15 examples", sh, srow, [9, 30, 24, 8, 36, 30, 16, 40, 22, 70, 70])
    text_sheet(wb, "Column guide", [
        "COLUMN GUIDE",
        "record ID - permanent name of the record: the source, then the source's own identifier. Use it to refer to a row.",
        "language (machine) - rule-based guess from script and common words. Least reliable on short, Romanized and mixed text.",
        "words - letter-runs separated by spaces/punctuation; Hindi vowel signs stay attached. Not comparable across languages.",
        "page type (machine-assigned) - what the automated search judged the page to be after opening it. Not a human judgement.",
        "status - 'candidate' means it passed automatic checks only. 'kept separately' pages are contrast material, never mixed in.",
        "author / commenter / video group - scrambled labels that let you see when two rows share an author or a video. No names are stored.",
        "remarks - 'extra page from the same site or author' (more than 3 pages from one site or more than 1 from one author); 'very long'; 'lines of site text removed from the end'; 'repost' or 'unclear whether original'.",
        "[HANDLE] [URL] [EMAIL] [PHONE] [NAME] - details masked in the cleaned text."])
    wb.save(OUT / "2 Clean data - Blogs and YouTube.xlsx")

    wb = new_wb()
    early = sum("focus_early" in r["flags"] for r in red)
    text_sheet(wb, "READ ME", [
        "EXPLORATORY MATERIAL ONLY - REDDIT POSTS AND THE YORK GRIEF SURVEY (release v1)", f"Built {BUILT}.", "", caution, "",
        "Neither source belongs in a Hindi-English comparison. Both are English. Use them to develop and test codebook categories.", "",
        f"Reddit posts - {len(red):,} posts (2019-2021) from r/offmychest, r/trueoffmychest, r/confession, r/confessions that passed the bereavement keyword screen. {early:,} mention the death in the title or first 500 characters ('death mentioned early' = yes): in a small machine-read sample 8 of 18 of those were really about a bereavement, against 1 of 12 of the others. No usernames, no replies. No suitable Hindi material was found in the Reddit sources searched (which does not prove none exists).",
        "Source: SocialGrep/one-million-reddit-confessions on Hugging Face, CC BY 4.0 (label given by its compiler). Posts shown as removed or deleted in the source were never kept.", "",
        f"York survey - {len(york)} UK respondents' written answers from 'Experiences of Grief: A Phenomenological Survey, 2020' (Millar, Ratcliffe, Richardson, Hughes; UK Data Service ReShare 856067, doi:10.5255/UKDA-SN-856067, CC BY 4.0; participants agreed to public release).",
        "CORRECTION: an earlier note said respondents wrote a median of about 1,900 words. That was wrong - the spreadsheet has a column repeating every answer and it had been counted. The verified median is 639 words (197,899 words, 235 respondents). Only the 40 answer columns are used here.", "",
        "Texts longer than Excel allows in one cell continue in 'text (part 2)'."])
    rh = ["record ID", "community", "posted (UTC)", "link", "words", "death mentioned early", "status", "text (part 1)", "text (part 2)"]
    rrow = [[r["record_id"], r["genre"].split()[1], datetime.fromtimestamp(int(r["published_at"]), timezone.utc).strftime("%Y-%m-%d"), r["source_url"], r["n_words"],
             "yes" if "focus_early" in r["flags"] else "no", "passed keyword screen (not human-checked)"] + parts(r["text_clean"], 2) for r in red]
    sheet(wb, "Reddit posts", rh, rrow, [30, 18, 13, 40, 8, 12, 34, 110, 30])
    yh2 = ["record ID", "person who died (survey field)", "past or current grief (survey field)", "nationality (self-reported)", "words (answers only)", "answers (question codes in brackets)"]
    yrow2 = [[r["record_id"], (r.get("extra") or {}).get("person_who_died"), (r.get("extra") or {}).get("past_or_present"), (r.get("extra") or {}).get("nationality_self_reported"),
              word_count(re.sub(r"\[[AB]\d\d\]", " ", r["text_clean"])), parts(r["text_clean"], 1)[0]] for r in york]
    sheet(wb, "York survey", yh2, yrow2, [18, 22, 40, 18, 12, 120])
    wb.save(OUT / "3 Clean data - EXPLORATORY Reddit and York.xlsx")

    wb = new_wb()
    rev_cols = [c for c in a_sheet[0] if c.startswith(("r1_", "r2_", "adjudicated_")) or c == "review_date"]
    def rev_name(c):
        c = re.sub(r"^r([12])_", r"Reviewer \1: ", c).replace("adjudicated_", "Agreed after discussion: ")
        return c.replace("_", " ")
    ah = ["ID (do not change)", "kind of text", "context", "words", "text (part 1)", "text (part 2)", "text (part 3)"] + [rev_name(c) for c in rev_cols]
    arow = []
    for s, k in zip(a_sheet, a_key):
        assert s["annotation_id"] == k["annotation_id"]
        arow.append([s["annotation_id"], plain(s["source_genre"]).replace(" (reused dataset)", ""), s["thread_context"], int(s["n_words_machine"])] + parts(rec[k["record_id"]]["text_clean"], 3) + [""] * len(rev_cols))
    sheet(wb, "Annotation sheet", ah, arow, [16, 30, 34, 8, 110, 20, 20] + [18] * len(rev_cols))
    src = ROOT / "release" / "v1_src" / "reviewer"
    instr = (src / "REVIEWER_INSTRUCTIONS.md").read_text(encoding="utf-8")
    instr = re.sub(r"## What you have.*?## Ground rules", "## What you have\nThis workbook. Sheet 'Annotation sheet' has one row per text; long texts continue in 'text (part 2)' and '(part 3)'. "
                   "Sheet 'Codebook status' lists the columns you fill and the state of their definitions.\n\n## Ground rules", instr, flags=re.S)
    text_sheet(wb, "Instructions", [unmark(l) for l in instr.splitlines()])
    cb = (src / "CODEBOOK_STATUS.md").read_text(encoding="utf-8").splitlines()
    table = [[unmark(c.strip()) for c in l.strip().strip("|").split("|")] for l in cb if l.startswith("|") and not set(l.strip()) <= set("|-: ")]
    prose = [unmark(l) for l in cb if l.strip() and not l.startswith("|")]
    ws = sheet(wb, "Codebook status", table[0], table[1:], [30, 60, 70, 60], note=" ".join(prose[:3]))
    for i, line in enumerate(prose[3:], ws.max_row + 2):
        ws.cell(row=i, column=1, value=line).font = ARIAL
    for row in ws.iter_rows(min_row=4, max_row=3 + len(table)):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.move_sheet("Instructions", offset=-1)
    wb.save(OUT / "4 Annotation sheet - BLIND for reviewers.xlsx")

    wb = new_wb()
    text_sheet(wb, "READ ME", ["BLOG FULL TEXTS - PRIVATE TO THE PROJECT", f"Built {BUILT}.", "",
                               "These are other people's copyrighted writing about real losses, copied for our research only. Keep this file between the two of you. Do not post it, forward it, or paste texts into online tools. What may be published later: web addresses, record IDs, our codes, and short quotations.",
                               "", "", caution,
                               "Cleaning kept the original script, punctuation, emojis, pronouns and kin terms. Links, e-mails, handles and phone numbers are masked; some site text at the end of pages was stripped by rule, so poems, quoted verses or leftover site text may remain."])
    fh = ["record ID", "web address", "title", "language (machine)", "words", "page type (machine-assigned)", "status", "text (part 1)", "text (part 2)", "text (part 3)"]
    frow = [[r["record_id"], r["source_url"], (r.get("extra") or {}).get("title"), LANG.get(r["lang_provisional"]), r["n_words"],
             r["message_type_provisional"].replace("_", " "), status(r)] + parts(r["text_clean"], 3) for r in blogs]
    sheet(wb, "Blog full texts", fh, frow, [22, 45, 40, 18, 8, 22, 34, 110, 20, 20])
    wb.save(OUT / "5 Blog full texts - PRIVATE do not forward.xlsx")

    wb = new_wb()
    text_sheet(wb, "READ ME", ["DO NOT OPEN UNTIL INDEPENDENT ANNOTATION IS FINISHED", "",
                               "This links every ID in the annotation sheet to its source, web address, machine language label, machine-assigned type, sampling group, and author/video groups. "
                               "Seeing it before or during annotation would prime you; showing it to reviewers would spoil their independence. It is in a separate folder only as a precaution - nothing technical stops anyone opening it.",
                               "Needed afterwards to: compare human and machine labels, keep one author's texts together when splitting data, and trace rows to sources.",
                               "Each ID is permanent and will never be renumbered, even if rows are added later. Sheet 'Old to new IDs' lists the earlier numbering for reference."])
    KH = {"annotation_id": "ID", "record_id": "record ID", "source": "source", "stratum": "sampling group", "sampling_note": "how it was sampled",
          "lang_provisional_machine": "language (machine)", "message_type_provisional": "type (machine)", "message_type_source": "who assigned the type",
          "eligibility_status_machine": "status (machine)", "thread_id": "site or video group", "author_group": "author group (scrambled)",
          "duplicate_group": "duplicate group", "flags": "remarks", "source_url": "web address"}
    keep_raw = {"annotation_id", "record_id", "thread_id", "author_group", "duplicate_group", "source_url"}
    sheet(wb, "Machine key", [KH[h] for h in a_key[0]], [[v if h in keep_raw else plain(v).replace("ytcv", "YTCommentVerse").replace("hinemo9", "HinEmo-9") for h, v in k.items()] for k in a_key],
          [16, 40, 18, 28, 50, 20, 26, 34, 30, 34, 20, 40, 26, 50])
    mp = list(csv.DictReader((ROOT / "data" / "annotation" / "id_mapping_v1.csv").open(encoding="utf-8")))
    sheet(wb, "Old to new IDs", ["earlier number", "permanent ID", "record ID"], [list(m.values()) for m in mp], [16, 16, 44])
    wb.save(OUT / "DO NOT OPEN until annotation is finished" / "Annotation machine key.xlsx")

    q = Counter(k["lang_provisional_machine"] for k in a_key)
    qs = lambda src, l: sum(1 for k in a_key if k["source"] in src and k["lang_provisional_machine"] == l)
    hs = q["devanagari_hindi"] + q["romanized_hindi"] + q["mixed_hi_en"]
    wd = lambda l: sum(r["n_words"] for r in el if r["lang_provisional"] == l)
    nd = lambda l: sum(1 for r in el if r["lang_provisional"] == l)
    core = lambda l: [r for r in el if r["lang_provisional"] == l and not {"over_cap", "length_outlier_gt5000w"} & set(r["flags"])]
    yl = Counter(r["lang_provisional"] for r in ytcv)
    hl = Counter(r["lang_provisional"] for r in hin)
    summary = f"""HINDI-ENGLISH GRIEF PROJECT - WRITTEN SUMMARY FOR COLLABORATORS
Release v1. Written {BUILT}. Data frozen at {CUTOFF}.

READ THIS FIRST
This folder holds texts we collected and cleaned by machine, so that the two of us can
pilot a codebook. It is NOT a validated corpus, NOT a representative sample and NOT a
finished study. No human has checked a single row. Every label (language, type of text,
"candidate") came from rules, a keyword screen, or machine reading of pages.

FILES IN THIS FOLDER
1 SUMMARY - read this first.txt                    this file
2 Clean data - Blogs and YouTube.xlsx              main material: blogs (details + first 300 characters) and YouTube comments
3 Clean data - EXPLORATORY Reddit and York.xlsx    English-only side material for developing codes
4 Annotation sheet - BLIND for reviewers.xlsx      301 texts with empty reviewer columns + instructions
5 Blog full texts - PRIVATE do not forward.xlsx    full blog texts - copyrighted, keep between us
DO NOT OPEN until annotation is finished (folder)  machine labels for the 301 rows - please do not look yet

WHAT WAS COLLECTED
1. Blogs and personal essays - OUR MAIN MATERIAL
   {len(blogs)} web pages found by search and opened one at a time.
   {len(el)} look like first-person writing about the death of someone the writer knew:
      Hindi (Devanagari)  {nd('devanagari_hindi')} pages, {wd('devanagari_hindi'):,} words
      English             {nd('english')} pages, {wd('english'):,} words
      1 further page is heavily mixed Hindi-English (the machine calls it Romanized Hindi)
   {len(blogs) - len(el)} are contrast pages kept apart on purpose: poems, tributes to public figures,
   advice, condolence templates, translations, obituaries.
   Leaving out five very long Hindi memoirs (over 5,000 words each) and extra pages beyond
   three per site or one per author: {len(core('devanagari_hindi'))} Hindi pages ({sum(r['n_words'] for r in core('devanagari_hindi')):,} words) and
   {len(core('english'))} English pages ({sum(r['n_words'] for r in core('english')):,} words). Our goal was about 26 pages and 30,000 words per
   language, so both are met - provisionally, until humans confirm what each page is.
   Word counts are not comparable across the two languages.

2. YouTube comments - SECONDARY, SHORT TEXTS
   {len(ytcv):,} comments from an open research dataset (YTCommentVerse) + {len(hin)} from another (HinEmo-9).
   They passed a bereavement keyword screen. Machine language labels:
      YTCommentVerse: English {yl['english']}, Hindi Romanized {yl['romanized_hindi']}, Hindi Devanagari {yl['devanagari_hindi']}, mixed {yl['mixed_hi_en']}, uncertain {yl['uncertain']}
      HinEmo-9:       English {hl['english']}, Hindi Romanized {hl['romanized_hindi']}, Hindi Devanagari {hl['devanagari_hindi']}
   Only 22.2% of the big dataset was scanned: its download was cut off by the server and we
   froze the data there. Nothing in this folder waits on the rest.
   In the rows read by machine, about half were real first-person losses; the rest
   included TV-serial plot talk.

3. Reddit posts - EXPLORATORY ONLY, ENGLISH
   {len(red):,} posts from four "confession" communities, 2019-2021. {early:,} mention the death
   early in the post. No usernames, no replies. We found no suitable Hindi bereavement
   material in the Reddit sources we searched; that does not prove none exists.

4. York grief survey - EXPLORATORY ONLY, ENGLISH, UK
   {len(york)} people's written survey answers about a bereavement (openly licensed, participants
   agreed to publication). Useful for developing codes; not comparable with online writing.

5. Seed examples
   15 texts read in full at the start to test feasibility: 8 English, 6 Hindi (Devanagari),
   1 Hindi (Romanized). They are kept out of the annotation sheet.

NOT COLLECTED
   X / Twitter: 0 posts. Directly from YouTube: 0 comments. Reddit's researcher programme: not applied for.

THE ANNOTATION SHEET (workbook 4)
   301 texts, each with a permanent ID that will never be renumbered; reviewer columns completely empty.
   Because we have not agreed a codebook, it is ready for CODEBOOK PILOTING, not yet for
   formal independent annotation. Machine language labels of the 301 rows:

                        English  Hindi-Deva  Hindi-Roman  Mixed  Uncertain  Total
   Blogs                  {qs(['blogs'],'english'):>5}  {qs(['blogs'],'devanagari_hindi'):>10}  {qs(['blogs'],'romanized_hindi'):>11}  {qs(['blogs'],'mixed_hi_en'):>5}  {qs(['blogs'],'uncertain'):>9}  {sum(1 for k in a_key if k['source']=='blogs'):>5}
   YouTube comments       {qs(['ytcv','hinemo9'],'english'):>5}  {qs(['ytcv','hinemo9'],'devanagari_hindi'):>10}  {qs(['ytcv','hinemo9'],'romanized_hindi'):>11}  {qs(['ytcv','hinemo9'],'mixed_hi_en'):>5}  {qs(['ytcv','hinemo9'],'uncertain'):>9}  {sum(1 for k in a_key if k['source'] in ('ytcv','hinemo9')):>5}
   Reddit (exploratory)   {qs(['reddit_confessions'],'english'):>5}  {0:>10}  {0:>11}  {0:>5}  {0:>9}  {sum(1 for k in a_key if k['source']=='reddit_confessions'):>5}
   York (exploratory)     {qs(['york_survey'],'english'):>5}  {0:>10}  {0:>11}  {0:>5}  {0:>9}  {sum(1 for k in a_key if k['source']=='york_survey'):>5}
   ALL                    {q['english']:>5}  {q['devanagari_hindi']:>10}  {q['romanized_hindi']:>11}  {q['mixed_hi_en']:>5}  {q['uncertain']:>9}  {len(a_key):>5}

   The imbalance is real: {hs} rows ({hs / len(a_key):.1%}) are Hindi in any form, {q['english']} ({q['english'] / len(a_key):.1%}) are English.
   Evening it out would have meant adding shayari, condolence templates, celebrity tributes,
   translated pages or TV-serial chatter, none of which is personal bereavement writing.
   The sheet was sampled to cover sources, so it says nothing about how common anything is.
   Twenty more Hindi YouTube comments turned up after the sheet was frozen; they are in
   workbook 2 and can be added later without changing any existing ID.

WHY THE YOUTUBE COMMENTS CANNOT SUPPORT A "SAME VIDEO, TWO LANGUAGES" CLAIM
   We wanted Hindi and English comments under videos we chose, with titles, dates and
   replies. This dataset scrambles video IDs, has no dates or replies, its videos were picked
   by its authors, and almost all Hindi comments are in Roman script from a few song videos
   (about half of the first 45 came from one grandmother-song video).

THE BIGGEST CAUTION FOR OUR COMPARISON
   Hindi pages are mostly 2007-2017 Blogspot death-anniversary posts and literary memoirs.
   English pages are mostly 2015-2026 news-site personal essays. Any difference we see
   between languages could be a difference of genre or period instead.
   We never treat writing in Hindi or English as evidence of nationality, religion or culture.

A CORRECTION
   An earlier note said York respondents wrote a median of about 1,900 words. Wrong: the
   spreadsheet repeats every answer in an extra column and it was counted. Verified median:
   639 words (197,899 words, 235 respondents). Everything here uses the corrected figure.

WHAT WE DO NOT HAVE (checked on official pages, 19 Sep 2026)
   YouTube access key (free): needed for the chosen-video pilot. The collection tool for it is
     prepared but has never been run with a key, so it is untested. YouTube's rules let an app keep
     comment data it pulls for at most 30 days before refreshing or deleting it; its researcher
     programme keeps that rule until data are fixed for final analysis and forbids passing data
     to third parties. Those rules cover data pulled directly from YouTube, not the open datasets we used.
   X / Twitter: $0.005 per post read on the public price page; search pricing and minimum purchase
     are not shown there. Every retrieved post costs money whether or not it is a grief post, so
     20,000 usable posts could cost $200 (if half are usable) to $2,000 (if 1 in 20) - assumed
     rates, not measured. Nothing has been spent and no budget has been authorised.
   Reddit researcher programme (optional): needs an accredited-university affiliation, ethics
     approval, an institutional sponsor; no passing data on; one year at most. It would not give us Hindi.

DECISIONS FOR THE TWO OF US
   1. Written definitions, with examples in both languages, for: personal distress, collective
      coping, continuing bonds, responsibility to survivors, indirect reference to death, addressee.
   2. For blogs: analyse the whole page or only the writer's own prose (pages embed poems,
      sometimes written by the person who died)?
   3. How to handle the genre/period difference: model it, search again for matching genres,
      or narrow the comparison.
   4. Keep Romanized and mixed text as separate categories, or fold them into Hindi?
   5. Short texts: wait for a YouTube key and build the matched pilot, or treat the current
      comments as a case study only?
   6. Who reviews, and when the machine-label folder may be opened.
   7. What could ever be published: IDs, web addresses, codes, short quotations - not full texts.

CREDITS REQUIRED WHEN REUSING
   YTCommentVerse (Zenodo doi:10.5281/zenodo.15678816, CC BY 4.0); HinEmo-9 (Zenodo
   doi:10.5281/zenodo.22041925, CC BY 4.0); SocialGrep one-million-reddit-confessions (Hugging
   Face, CC BY 4.0); Experiences of Grief survey, Millar, Ratcliffe, Richardson & Hughes
   (doi:10.5255/UKDA-SN-856067, CC BY 4.0). Blog pages belong to their authors.

"""
    (OUT / "1 SUMMARY - read this first.txt").write_text(summary, encoding="utf-8")

    z = OUT.with_suffix(".zip")
    if z.exists():
        z.unlink()
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for p in sorted(OUT.rglob("*")):
            if p.is_file() and p.name != ".DS_Store":
                zf.write(p, (Path(OUT.name) / p.relative_to(OUT)).as_posix())
    for p in sorted(OUT.rglob("*")):
        if p.is_file():
            print(f"{p.stat().st_size / 1e6:8.2f} MB  {p.relative_to(OUT)}")
    print(f"{z.stat().st_size / 1e6:8.2f} MB  ZIP {z.name}")


if __name__ == "__main__":
    main()
