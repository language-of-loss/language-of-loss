import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "release" / "Hindi and English data files"
ARIAL, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
FILL = PatternFill("solid", fgColor="DDE6F0")
LANG = {"english": "English", "devanagari_hindi": "Hindi (Devanagari script)", "romanized_hindi": "Hindi (Roman script)",
        "mixed_hi_en": "Mixed Hindi-English", "uncertain": "Unclear"}
HINDI = {"devanagari_hindi", "romanized_hindi", "mixed_hi_en"}
PART = 32000


def jl(name):
    return [json.loads(l) for l in (ROOT / "data" / "clean" / f"{name}.jsonl").open(encoding="utf-8")]


def parts(text, n=3):
    text = ILLEGAL_CHARACTERS_RE.sub("", text or "")
    chunks = [text[i:i + PART] for i in range(0, len(text), PART)] or [""]
    return chunks + [""] * (n - len(chunks))


def add_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    for j, h in enumerate(header, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = BOLD, FILL
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"
            c.font = ARIAL
    ws.freeze_panes = "B2"
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def about(wb, lines):
    ws = wb.create_sheet("About this file", 0)
    ws.column_dimensions["A"].width = 125
    for i, line in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.font = BOLD if i == 1 else ARIAL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def blog_rows(recs):
    kind = lambda r: "personal account of a loss (candidate)" if r["eligibility_status"] == "provisionally_eligible" \
        else "comparison page: " + r["message_type_provisional"].replace("_", " ")
    recs = sorted(recs, key=lambda r: (r["eligibility_status"] != "provisionally_eligible", r["record_id"]))
    return [[r["record_id"], kind(r), LANG[r["lang_provisional"]], r["n_words"], r["published_at"], (r.get("extra") or {}).get("title"),
             (r.get("extra") or {}).get("deceased_relation_as_stated"), r["source_url"]] + parts(r["text_clean"]) for r in recs]


def yt_rows(recs):
    return [[r["record_id"], LANG[r["lang_provisional"]], r["n_words"], r["thread_id"] or "not available", r["text_clean"]] for r in recs]


BH = ["ID", "what it is", "language", "words", "published", "title", "who died (as the page says)", "web address", "text (part 1)", "text (part 2)", "text (part 3)"]
BW = [22, 36, 22, 8, 12, 40, 24, 45, 110, 20, 20]
YH = ["ID", "language", "words", "video (scrambled label; same label = same video)", "comment text"]
YW = [40, 24, 8, 36, 120]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    blogs, yt = jl("blogs"), jl("ytcv") + jl("hinemo9")
    note = ("Language, 'what it is' and every other label were assigned by rules or by machine reading. No person has checked any row. "
            "Texts are cleaned: links, e-mail addresses, user handles and phone numbers are masked; original script, spelling, punctuation and emojis are kept. "
            "Very long texts continue in 'text (part 2)' and 'text (part 3)'.")
    credit = ("YouTube comments come from two open research datasets - YTCommentVerse (doi:10.5281/zenodo.15678816) and HinEmo-9 (doi:10.5281/zenodo.22041925), "
              "both CC BY 4.0, credit required. Blog pages belong to their authors: for our research only, please do not post or forward them.")
    out = {}
    for name, keep_b, keep_y in (("Hindi", lambda r: r["lang_provisional"] in HINDI, lambda r: r["lang_provisional"] in HINDI),
                                 ("English", lambda r: r["lang_provisional"] == "english", lambda r: r["lang_provisional"] in ("english", "uncertain"))):
        b, y = [r for r in blogs if keep_b(r)], [r for r in yt if keep_y(r)]
        nb = sum(1 for r in b if r["eligibility_status"] == "provisionally_eligible")
        wb = Workbook()
        wb.remove(wb.active)
        extra = ("Hindi here means Devanagari script, Roman script, and mixed Hindi-English; the 'language' column says which."
                 if name == "Hindi" else "Two very short comments whose language the rules could not decide are included and marked 'Unclear'.")
        about(wb, [f"{name} data - Hindi-English grief project", "",
                   f"Blogs: {len(b)} pages - {nb} candidate personal accounts of a loss, listed first, then {len(b) - nb} comparison pages (poems, tributes, advice and similar) kept for contrast.",
                   f"YouTube comments: {len(y)} comments that mention a bereavement.", extra, "", note, "", credit])
        add_sheet(wb, f"{name} blogs", BH, blog_rows(b), BW)
        add_sheet(wb, f"{name} YouTube comments", YH, yt_rows(y), YW)
        path = OUT / f"{name} data.xlsx"
        wb.save(path)
        out[name] = (len(b), nb, len(y), sum(r["n_words"] for r in b if r["eligibility_status"] == "provisionally_eligible"), path.stat().st_size)
    print(out, "| blogs total", len(blogs), "| youtube total", len(yt))


if __name__ == "__main__":
    sys.exit(main())
